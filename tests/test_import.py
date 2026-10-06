from pathlib import Path

import pymupdf
from docx import Document

from qualitative_themes.importers.extract import extract_document
from qualitative_themes.storage.database import create_project, connect
from qualitative_themes.storage.documents import get_document, import_document, list_documents


def test_import_txt_exact_bytes_and_duplicate(tmp_path):
    project = create_project(tmp_path / "study", "Study")
    src = tmp_path / "notes.txt"
    text = "Café — first thought.\r\n\r\nThe same phrase. The same phrase.\n"
    src.write_bytes(text.encode("utf-8"))
    outcome = import_document(project.directory, src)
    assert outcome.status == "imported"
    name, saved, spans = get_document(project.directory, outcome.document_id)
    assert name == src.name
    assert saved == text
    assert spans == [(None, None, 0, len(text))]
    assert import_document(project.directory, src).status == "duplicate"
    assert len(list_documents(project.directory)) == 1


def test_pdf_page_spans_and_reopen(tmp_path):
    project = create_project(tmp_path / "study", "Study")
    src = tmp_path / "two-pages.pdf"
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "First page: access to care.")
    pdf.new_page().insert_text((72, 72), "Second page: transport barriers.")
    pdf.save(src)
    pdf.close()
    outcome = import_document(project.directory, src)
    assert outcome.status == "imported"
    _, saved, spans = get_document(project.directory, outcome.document_id)
    assert [x[0] for x in spans] == [1, 2]
    assert "First page" in saved[spans[0][2]:spans[0][3]]
    assert "Second page" in saved[spans[1][2]:spans[1][3]]
    assert "First page" not in saved[spans[1][2]:spans[1][3]]


def test_docx_paragraphs_table_and_unicode(tmp_path):
    src = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("Résumé: services")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Site A"
    table.cell(0, 1).text = "Cost concerns"
    doc.add_paragraph("Final paragraph.")
    doc.save(src)
    extracted = extract_document(src)
    assert "Résumé: services" in extracted.text
    assert "Site A\tCost concerns" in extracted.text
    assert extracted.text.index("Cost concerns") < extracted.text.index("Final paragraph")
    assert [s.paragraph_number for s in extracted.spans] == [1, 2, 3]
    assert all(extracted.text[s.start_char:s.end_char].strip() for s in extracted.spans)


def test_empty_pdf_fails_without_project_row_or_source_copy(tmp_path):
    project = create_project(tmp_path / "study", "Study")
    src = tmp_path / "scan.pdf"
    pdf = pymupdf.open()
    pdf.new_page()
    pdf.save(src)
    pdf.close()
    outcome = import_document(project.directory, src)
    assert outcome.status == "failed"
    assert "OCR" in outcome.message
    assert list_documents(project.directory) == []
    assert list((project.directory / "sources").iterdir()) == []


def test_same_filename_different_content_is_not_overwritten(tmp_path):
    project = create_project(tmp_path / "study", "Study")
    a, b = tmp_path / "a", tmp_path / "b"
    a.mkdir(); b.mkdir()
    (a / "notes.txt").write_text("Original note.")
    (b / "notes.txt").write_text("Revised note.")
    one = import_document(project.directory, a / "notes.txt")
    two = import_document(project.directory, b / "notes.txt")
    assert one.status == two.status == "imported"
    with connect(project.directory) as conn:
        paths = [row[0] for row in conn.execute("SELECT source_path FROM documents")]
    assert len(set(paths)) == 2
    assert all((project.directory / p).is_file() for p in paths)
