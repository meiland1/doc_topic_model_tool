"""Extract an immutable, displayable text snapshot with source spans.

Offsets refer to Python character indices in ``ExtractedDocument.text``.
Page and paragraph separators are outside the associated source spans.
"""
from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class SourceSpan:
    start_char: int
    end_char: int
    page_number: int | None = None
    paragraph_number: int | None = None


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    spans: tuple[SourceSpan, ...]
    warnings: tuple[str, ...] = ()
    encoding: str | None = None


class _TextBuilder:
    def __init__(self):
        self.parts: list[str] = []
        self.length = 0
        self.spans: list[SourceSpan] = []

    def add(self, text: str, separator: str, *, page: int | None = None,
            paragraph: int | None = None):
        if not text.strip():
            return
        if self.parts:
            self.parts.append(separator)
            self.length += len(separator)
        start = self.length
        self.parts.append(text)
        self.length += len(text)
        self.spans.append(SourceSpan(start, self.length, page, paragraph))

    def finish(self, **kwargs) -> ExtractedDocument:
        return ExtractedDocument("".join(self.parts), tuple(self.spans), **kwargs)


def _pdf(path: Path) -> ExtractedDocument:
    import pymupdf

    builder = _TextBuilder()
    with pymupdf.open(path) as pdf:
        if pdf.needs_pass:
            raise ValueError("This PDF is password protected. Unlock it before importing.")
        for number, page in enumerate(pdf, start=1):
            # This is a reading-order approximation, not a map of PDF glyph coordinates.
            builder.add(page.get_text("text", sort=True), "\n\f\n", page=number)
        if not builder.spans:
            raise ValueError("No readable text was found. This PDF may need OCR.")
    return builder.finish(warnings=(
        "PDF text order may differ from the visual page layout; check the preview, especially for columns or tables.",
    ))


def _docx(path: Path) -> ExtractedDocument:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    builder = _TextBuilder()
    paragraph_no = 0
    # iter_inner_content preserves top-level paragraph/table order in python-docx 1.1+.
    for item in doc.iter_inner_content():
        if isinstance(item, Paragraph):
            paragraph_no += 1
            builder.add(item.text, "\n\n", paragraph=paragraph_no)
        elif isinstance(item, Table):
            for row in item.rows:
                paragraph_no += 1
                builder.add("\t".join(cell.text for cell in row.cells), "\n", paragraph=paragraph_no)
    if not builder.spans:
        raise ValueError("No readable body text was found in this Word document.")
    return builder.finish(warnings=(
        "Headers, footers, text boxes, and embedded images are not included in this text preview.",
    ))


def _txt(path: Path) -> ExtractedDocument:
    data = path.read_bytes()
    warning = ()
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        text, encoding = data.decode("utf-16"), "utf-16"
    else:
        try:
            text, encoding = data.decode("utf-8-sig"), "utf-8-sig"
        except UnicodeDecodeError:
            text, encoding = data.decode("cp1252"), "cp1252"
            warning = ("Text was read using Windows-1252 encoding; check accented characters in the preview.",)
    if not text.strip():
        raise ValueError("This text file is empty.")
    # No normalization: all offsets point to the exact saved decoded text.
    return ExtractedDocument(text, (SourceSpan(0, len(text)),), warning, encoding)


def extract_document(path: Path) -> ExtractedDocument:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        result = _pdf(path)
    elif ext == ".docx":
        result = _docx(path)
    elif ext == ".txt":
        result = _txt(path)
    else:
        raise ValueError(f"Unsupported file type: {ext or '(none)'}. Choose PDF, DOCX, or TXT.")
    for span in result.spans:
        assert 0 <= span.start_char < span.end_char <= len(result.text)
    return result
